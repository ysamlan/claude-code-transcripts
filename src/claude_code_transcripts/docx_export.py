"""DOCX export for Claude Code transcripts."""

import json
import re

import click

from . import COMMIT_PATTERN


def _add_hyperlink(
    paragraph, text, url, font_name=None, font_size=None, color_rgb=None
):
    """Add a clickable hyperlink run to a paragraph.

    python-docx doesn't have built-in hyperlink support, so we build the XML directly.
    """
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)

    new_run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")

    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)

    c = OxmlElement("w:color")
    if color_rgb:
        c.set(qn("w:val"), f"{color_rgb[0]:02X}{color_rgb[1]:02X}{color_rgb[2]:02X}")
    else:
        c.set(qn("w:val"), "0563C1")
    rPr.append(c)

    if font_name:
        rFonts = OxmlElement("w:rFonts")
        rFonts.set(qn("w:ascii"), font_name)
        rFonts.set(qn("w:hAnsi"), font_name)
        rPr.append(rFonts)

    if font_size:
        sz = OxmlElement("w:sz")
        sz.set(qn("w:val"), str(font_size * 2))  # half-points
        rPr.append(sz)

    new_run.append(rPr)

    t = OxmlElement("w:t")
    t.text = text
    new_run.append(t)

    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def _set_paragraph_shading(paragraph, hex_color):
    """Set background color on a paragraph via XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def _set_paragraph_border(paragraph, color, side="left", size=12):
    """Set a single border on a paragraph via XML."""
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    border = OxmlElement(f"w:{side}")
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), str(size))
    border.set(qn("w:space"), "4")
    border.set(qn("w:color"), color)
    pBdr.append(border)


def _add_markdown_runs(paragraph, text, Pt, RGBColor, base_size=10):
    """Parse markdown inline formatting and add styled runs to a paragraph.

    Handles: **bold**, *italic*, `code`, [text](url).
    """
    # Tokenize the text into segments with formatting
    # Pattern matches: **bold**, *italic*, `code`, [text](url)
    pattern = re.compile(
        r"(\*\*(.+?)\*\*)"  # bold
        r"|(\*(.+?)\*)"  # italic
        r"|(`([^`]+)`)"  # inline code
        r"|(\[([^\]]+)\]\(([^)]+)\))"  # links
    )

    pos = 0
    for m in pattern.finditer(text):
        # Add text before this match
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            run.font.size = Pt(base_size)

        if m.group(2):  # bold
            run = paragraph.add_run(m.group(2))
            run.bold = True
            run.font.size = Pt(base_size)
        elif m.group(4):  # italic
            run = paragraph.add_run(m.group(4))
            run.italic = True
            run.font.size = Pt(base_size)
        elif m.group(6):  # code
            run = paragraph.add_run(m.group(6))
            run.font.name = "Courier New"
            run.font.size = Pt(base_size - 1)
            run.font.color.rgb = RGBColor(0xC6, 0x28, 0x28)
        elif m.group(8):  # link
            _add_hyperlink(paragraph, m.group(8), m.group(9))

        pos = m.end()

    # Add remaining text
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        run.font.size = Pt(base_size)


def generate_docx(source, output_path, github_repo=None):
    """Generate a DOCX from a session file or session data dict.

    Args:
        source: Path to session file, or dict with 'loglines' key
        output_path: Path for the output DOCX file
        github_repo: Optional GitHub repo for commit links
    """
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor
    except ImportError:
        raise click.ClickException(
            "python-docx is required for DOCX export. "
            "Install with: uv pip install 'claude-code-transcripts[docx]'"
        )

    from .export_helpers import load_conversations, count_prompts_and_messages
    from . import extract_text_from_content, is_tool_result_message

    conversations, github_repo = load_conversations(source, github_repo)

    doc = Document()

    # Set default font
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(10)

    # Title
    doc.add_heading("Claude Code Transcript", level=1)

    # Stats
    prompt_count, message_count = count_prompts_and_messages(conversations)
    stats_para = doc.add_paragraph()
    stats_run = stats_para.add_run(
        f"{prompt_count} prompts \u00b7 {message_count} messages"
    )
    stats_run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
    stats_run.font.size = Pt(9)

    doc.add_paragraph()  # spacer

    for conv in conversations:
        for log_type, message_json, timestamp in conv["messages"]:
            if not message_json:
                continue
            try:
                message_data = json.loads(message_json)
            except json.JSONDecodeError:
                continue

            content = message_data.get("content", "")

            if log_type == "user":
                if is_tool_result_message(message_data):
                    _render_tool_results_docx(
                        doc, content, timestamp, github_repo, Pt, RGBColor
                    )
                else:
                    # User message
                    header = _add_role_header(
                        doc, "USER", timestamp, RGBColor(0x19, 0x76, 0xD2), Pt, RGBColor
                    )
                    _set_paragraph_shading(header, "E3F2FD")
                    _set_paragraph_border(header, "1976D2", "left", 12)
                    text = extract_text_from_content(content)
                    if text:
                        para = doc.add_paragraph(text)
                        _set_paragraph_shading(para, "E3F2FD")
                        _set_paragraph_border(para, "1976D2", "left", 12)

            elif log_type == "assistant":
                if isinstance(content, list):
                    for block in content:
                        if not isinstance(block, dict):
                            continue
                        block_type = block.get("type", "")

                        if block_type == "text":
                            text = block.get("text", "")
                            if text:
                                _add_role_header(
                                    doc,
                                    "ASSISTANT",
                                    timestamp,
                                    RGBColor(0x9E, 0x9E, 0x9E),
                                    Pt,
                                    RGBColor,
                                )
                                _add_markdown_text(doc, text, Pt, RGBColor)

                        elif block_type == "thinking":
                            text = block.get("thinking", "")
                            if text:
                                para = doc.add_paragraph()
                                run = para.add_run("Thinking: ")
                                run.font.color.rgb = RGBColor(0xF5, 0x7C, 0x00)
                                run.font.size = Pt(8)
                                run.bold = True
                                run = para.add_run(text[:500])
                                run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                                run.font.size = Pt(8)
                                run.italic = True
                                if len(text) > 500:
                                    run = para.add_run("...")
                                    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
                                    run.font.size = Pt(8)
                                _set_paragraph_shading(para, "FFF8E1")
                                _set_paragraph_border(para, "FFC107", "left", 12)

                        elif block_type == "image":
                            _add_image_block_docx(doc, block)

                        elif block_type == "tool_use":
                            tool_name = block.get("name", "Unknown")
                            tool_input = block.get("input", {})
                            _render_tool_use_docx(
                                doc, tool_name, tool_input, Pt, RGBColor
                            )

                        elif block_type == "tool_result":
                            result_content = block.get("content", "")
                            is_error = block.get("is_error", False)
                            _render_single_tool_result_docx(
                                doc,
                                result_content,
                                is_error,
                                github_repo,
                                Pt,
                                RGBColor,
                            )

    from pathlib import Path

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))


def _add_role_header(doc, role, timestamp, color, Pt, RGBColor):
    """Add a role header paragraph (USER, ASSISTANT, etc.). Returns the paragraph."""
    para = doc.add_paragraph()
    run = para.add_run(role)
    run.bold = True
    run.font.color.rgb = color
    run.font.size = Pt(8)
    if timestamp:
        run = para.add_run(f"  {timestamp}")
        run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
        run.font.size = Pt(7)
    return para


def _add_markdown_text(doc, text, Pt, RGBColor):
    """Add text with basic markdown formatting to the document."""
    lines = text.split("\n")
    current_block = []

    def flush_block():
        if current_block:
            joined = "\n".join(current_block)
            if joined.strip():
                para = doc.add_paragraph()
                _add_markdown_runs(para, joined, Pt, RGBColor)
            current_block.clear()

    in_code_block = False
    code_lines = []

    for line in lines:
        if line.startswith("```"):
            if in_code_block:
                # End code block
                in_code_block = False
                flush_block()
                para = doc.add_paragraph()
                code_text = "\n".join(code_lines)
                run = para.add_run(code_text)
                run.font.name = "Courier New"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(0xAE, 0xD5, 0x81)
                _set_paragraph_shading(para, "263238")
                code_lines.clear()
            else:
                # Start code block
                in_code_block = True
                flush_block()
        elif in_code_block:
            code_lines.append(line)
        elif line.startswith("# "):
            flush_block()
            doc.add_heading(line[2:], level=2)
        elif line.startswith("## "):
            flush_block()
            doc.add_heading(line[3:], level=3)
        elif line.startswith("### "):
            flush_block()
            doc.add_heading(line[4:], level=4)
        elif line.startswith("- ") or line.startswith("* "):
            flush_block()
            para = doc.add_paragraph(style="List Bullet")
            _add_markdown_runs(para, line[2:], Pt, RGBColor)
        else:
            current_block.append(line)

    # Flush any remaining code block
    if in_code_block and code_lines:
        para = doc.add_paragraph()
        code_text = "\n".join(code_lines)
        run = para.add_run(code_text)
        run.font.name = "Courier New"
        run.font.size = Pt(8)
        _set_paragraph_shading(para, "263238")

    flush_block()


def _render_tool_use_docx(doc, tool_name, tool_input, Pt, RGBColor):
    """Render a tool use block."""
    para = doc.add_paragraph()
    run = para.add_run(f"\u2699 {tool_name}")
    run.bold = True
    run.font.color.rgb = RGBColor(0x9C, 0x27, 0xB0)
    run.font.size = Pt(9)
    _set_paragraph_shading(para, "F3E5F5")
    _set_paragraph_border(para, "9C27B0", "left", 12)

    if tool_name == "Write":
        file_path = tool_input.get("file_path", "")
        if file_path:
            para = doc.add_paragraph()
            run = para.add_run(f"  \u2192 {file_path}")
            run.font.name = "Courier New"
            run.font.size = Pt(8)
        content = tool_input.get("content", "")
        if content:
            para = doc.add_paragraph()
            run = para.add_run(content[:1000])
            run.font.name = "Courier New"
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0xAE, 0xD5, 0x81)
            _set_paragraph_shading(para, "263238")
            if len(content) > 1000:
                run = para.add_run(f"\n... ({len(content)} chars total)")
                run.font.size = Pt(7)
                run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)

    elif tool_name == "Edit":
        file_path = tool_input.get("file_path", "")
        if file_path:
            para = doc.add_paragraph()
            run = para.add_run(f"  \u2192 {file_path}")
            run.font.name = "Courier New"
            run.font.size = Pt(8)
        old_str = tool_input.get("old_string", "")
        new_str = tool_input.get("new_string", "")
        if old_str:
            para = doc.add_paragraph()
            run = para.add_run(f"- {old_str[:500]}")
            run.font.name = "Courier New"
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)
        if new_str:
            para = doc.add_paragraph()
            run = para.add_run(f"+ {new_str[:500]}")
            run.font.name = "Courier New"
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0x1B, 0x5E, 0x20)

    elif tool_name == "Bash":
        command = tool_input.get("command", "")
        if command:
            para = doc.add_paragraph()
            run = para.add_run(f"$ {command}")
            run.font.name = "Courier New"
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0xAE, 0xD5, 0x81)
            _set_paragraph_shading(para, "263238")

    else:
        # Generic tool - show input as JSON
        description = tool_input.get("description", "")
        if description:
            para = doc.add_paragraph()
            run = para.add_run(description)
            run.italic = True
            run.font.size = Pt(8)
            run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
        display_input = {k: v for k, v in tool_input.items() if k != "description"}
        if display_input:
            input_text = json.dumps(display_input, indent=2, ensure_ascii=False)
            if len(input_text) > 500:
                input_text = input_text[:500] + "..."
            para = doc.add_paragraph()
            run = para.add_run(input_text)
            run.font.name = "Courier New"
            run.font.size = Pt(7)
            run.font.color.rgb = RGBColor(0xAE, 0xD5, 0x81)
            _set_paragraph_shading(para, "263238")


def _render_tool_results_docx(
    doc, content_blocks, timestamp, github_repo, Pt, RGBColor
):
    """Render tool result blocks from a user message."""
    header = _add_role_header(
        doc, "TOOL REPLY", timestamp, RGBColor(0xE6, 0x51, 0x00), Pt, RGBColor
    )
    _set_paragraph_shading(header, "E8F5E9")
    _set_paragraph_border(header, "4CAF50", "left", 12)
    for block in content_blocks:
        if not isinstance(block, dict):
            continue
        _render_single_tool_result_docx(
            doc,
            block.get("content", ""),
            block.get("is_error", False),
            github_repo,
            Pt,
            RGBColor,
        )


def _render_single_tool_result_docx(
    doc, result_content, is_error, github_repo, Pt, RGBColor
):
    """Render a single tool result."""
    bg_color = "FFEBEE" if is_error else "E8F5E9"
    border_color = "B71C1C" if is_error else "4CAF50"

    if isinstance(result_content, str):
        # Check for commits
        commits = list(COMMIT_PATTERN.finditer(result_content))
        if commits:
            # Render text between/around commits
            last_end = 0
            for match in commits:
                before = result_content[last_end : match.start()].strip()
                if before:
                    para = doc.add_paragraph()
                    run = para.add_run(before[:2000])
                    run.font.name = "Courier New"
                    run.font.size = Pt(7)
                    _set_paragraph_shading(para, bg_color)
                    _set_paragraph_border(para, border_color, "left", 12)
                commit_hash = match.group(1)
                commit_msg = match.group(2)
                para = doc.add_paragraph()
                if github_repo:
                    link = f"https://github.com/{github_repo}/commit/{commit_hash}"
                    _add_hyperlink(
                        para,
                        f"\U0001f4e6 {commit_hash[:7]}",
                        link,
                        font_name="Courier New",
                        font_size=9,
                        color_rgb=(0xE6, 0x51, 0x00),
                    )
                else:
                    run = para.add_run(f"\U0001f4e6 {commit_hash[:7]}")
                    run.font.name = "Courier New"
                    run.font.color.rgb = RGBColor(0xE6, 0x51, 0x00)
                    run.font.size = Pt(9)
                run = para.add_run(f" {commit_msg}")
                run.font.size = Pt(9)
                _set_paragraph_shading(para, bg_color)
                _set_paragraph_border(para, border_color, "left", 12)
                last_end = match.end()
            after = result_content[last_end:].strip()
            if after:
                para = doc.add_paragraph()
                run = para.add_run(after[:2000])
                run.font.name = "Courier New"
                run.font.size = Pt(7)
                _set_paragraph_shading(para, bg_color)
                _set_paragraph_border(para, border_color, "left", 12)
        else:
            text = result_content[:2000]
            if text.strip():
                para = doc.add_paragraph()
                run = para.add_run(text)
                run.font.name = "Courier New"
                run.font.size = Pt(7)
                if is_error:
                    run.font.color.rgb = RGBColor(0xB7, 0x1C, 0x1C)
                if len(result_content) > 2000:
                    run = para.add_run(f"\n... ({len(result_content)} chars total)")
                    run.font.size = Pt(7)
                    run.font.color.rgb = RGBColor(0x75, 0x75, 0x75)
                _set_paragraph_shading(para, bg_color)
                _set_paragraph_border(para, border_color, "left", 12)
    elif isinstance(result_content, list):
        for item in result_content:
            if not isinstance(item, dict):
                continue
            item_type = item.get("type", "")
            if item_type == "text":
                text = item.get("text", "")
                if text.strip():
                    para = doc.add_paragraph()
                    run = para.add_run(text[:2000])
                    run.font.name = "Courier New"
                    run.font.size = Pt(7)
                    _set_paragraph_shading(para, bg_color)
                    _set_paragraph_border(para, border_color, "left", 12)
            elif item_type == "image":
                _add_image_block_docx(doc, item)


def _add_image_block_docx(doc, block):
    """Add an inline base64 image to the DOCX."""
    import base64
    import io

    from docx.shared import Inches

    source = block.get("source", {})
    data = source.get("data", "")
    if not data:
        return

    try:
        image_bytes = base64.b64decode(data)
        image_stream = io.BytesIO(image_bytes)
        # Max width ~6.5 inches (letter page with 1" margins)
        doc.add_picture(image_stream, width=Inches(6.5))
    except Exception:
        para = doc.add_paragraph()
        run = para.add_run("[Image could not be rendered]")
        run.italic = True
