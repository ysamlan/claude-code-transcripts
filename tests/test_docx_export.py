"""Tests for DOCX export functionality."""

import json
import tempfile
from pathlib import Path

import pytest

from claude_code_transcripts.docx_export import generate_docx


@pytest.fixture
def output_dir():
    """Create a temporary output directory."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield Path(tmpdir)


@pytest.fixture
def sample_loglines():
    return [
        {
            "type": "user",
            "timestamp": "2025-01-01T10:00:00.000Z",
            "message": {"role": "user", "content": "Hello, write a function"},
        },
        {
            "type": "assistant",
            "timestamp": "2025-01-01T10:00:05.000Z",
            "message": {
                "role": "assistant",
                "content": [
                    {
                        "type": "text",
                        "text": "Here is the function:\n```python\ndef hello():\n    print('hello')\n```",
                    }
                ],
            },
        },
        {
            "type": "user",
            "timestamp": "2025-01-01T10:01:00.000Z",
            "message": {"role": "user", "content": "Now add a parameter"},
        },
        {
            "type": "assistant",
            "timestamp": "2025-01-01T10:01:05.000Z",
            "message": {
                "role": "assistant",
                "content": [
                    {
                        "type": "text",
                        "text": "Updated:\n```python\ndef hello(name):\n    print(f'hello {name}')\n```",
                    }
                ],
            },
        },
    ]


class TestGenerateDocx:
    """Tests for DOCX generation."""

    def test_generates_docx_file_from_path(self, output_dir, sample_loglines):
        """Test that generate_docx creates a DOCX file from a session file."""
        session_file = output_dir / "test_session.json"
        session_file.write_text(
            json.dumps({"loglines": sample_loglines}), encoding="utf-8"
        )

        docx_path = output_dir / "output.docx"
        generate_docx(session_file, docx_path)

        assert docx_path.exists()
        assert docx_path.stat().st_size > 0

    def test_generates_docx_from_dict(self, output_dir, sample_loglines):
        """Test that generate_docx accepts a dict with loglines."""
        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": sample_loglines}, docx_path)

        assert docx_path.exists()
        assert docx_path.stat().st_size > 0

    def test_docx_contains_user_messages(self, output_dir, sample_loglines):
        """Test that the DOCX contains user message text."""
        from docx import Document

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": sample_loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Hello, write a function" in full_text
        assert "Now add a parameter" in full_text

    def test_docx_contains_assistant_text(self, output_dir, sample_loglines):
        """Test that the DOCX contains assistant response text."""
        from docx import Document

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": sample_loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "def hello" in full_text

    def test_docx_contains_title(self, output_dir, sample_loglines):
        """Test that the DOCX has a title."""
        from docx import Document

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": sample_loglines}, docx_path)

        doc = Document(str(docx_path))
        # Title should be in first paragraph (heading)
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Claude Code Transcript" in full_text

    def test_docx_contains_stats(self, output_dir, sample_loglines):
        """Test that the DOCX contains prompt/message counts."""
        from docx import Document

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": sample_loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "2 prompts" in full_text
        assert "4 messages" in full_text

    def test_creates_parent_directories(self, output_dir, sample_loglines):
        """Test that generate_docx creates parent directories."""
        docx_path = output_dir / "subdir" / "nested" / "output.docx"
        generate_docx({"loglines": sample_loglines}, docx_path)

        assert docx_path.exists()

    def test_docx_with_tool_use(self, output_dir):
        """Test DOCX generation with tool use blocks."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Create a file"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "tool_use",
                            "name": "Write",
                            "id": "tool_1",
                            "input": {
                                "file_path": "/tmp/test.py",
                                "content": "print('hello')",
                            },
                        }
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Write" in full_text
        assert "/tmp/test.py" in full_text

    def test_docx_with_image_block(self, output_dir):
        """Test that generate_docx handles image blocks without crashing."""
        # Valid 2x2 red PNG
        red_pixel = "iVBORw0KGgoAAAANSUhEUgAAAAIAAAACCAIAAAD91JpzAAAAEElEQVR4nGP8zwACTGCSAQANHQEDgslx/wAAAABJRU5ErkJggg=="

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Show me an image"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "image",
                            "source": {
                                "type": "base64",
                                "media_type": "image/png",
                                "data": red_pixel,
                            },
                        },
                        {"type": "text", "text": "Here is the image."},
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        assert docx_path.exists()
        assert docx_path.stat().st_size > 0

    def test_docx_with_thinking_block(self, output_dir):
        """Test that thinking blocks render in the DOCX."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Think about this"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "thinking",
                            "thinking": "Let me consider the options carefully...",
                        },
                        {"type": "text", "text": "Here is my answer."},
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Thinking:" in full_text
        assert "consider the options" in full_text
        assert "Here is my answer" in full_text

    def test_docx_with_generic_tool(self, output_dir):
        """Test DOCX generation with a non-Write/Edit/Bash tool."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Search for something"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "tool_use",
                            "name": "Grep",
                            "id": "tool_1",
                            "input": {
                                "description": "Search for pattern",
                                "pattern": "def hello",
                                "path": "/src",
                            },
                        }
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Grep" in full_text
        assert "Search for pattern" in full_text

    def test_docx_with_commit_output(self, output_dir):
        """Test DOCX renders commit hashes from tool results."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Commit this"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "tool_use",
                            "name": "Bash",
                            "id": "tool_1",
                            "input": {"command": "git commit -m 'test'"},
                        }
                    ],
                },
            },
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:10.000Z",
                "message": {
                    "role": "user",
                    "content": [
                        {
                            "type": "tool_result",
                            "tool_use_id": "tool_1",
                            "content": "[main abc1234] Add test feature\n 1 file changed",
                        }
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "abc1234" in full_text
        assert "Add test feature" in full_text

    def test_docx_with_error_tool_result(self, output_dir):
        """Test DOCX renders error tool results."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Run this"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "tool_use",
                            "name": "Bash",
                            "id": "tool_1",
                            "input": {"command": "exit 1"},
                        }
                    ],
                },
            },
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:10.000Z",
                "message": {
                    "role": "user",
                    "content": [
                        {
                            "type": "tool_result",
                            "tool_use_id": "tool_1",
                            "content": "Error: command failed with exit code 1",
                            "is_error": True,
                        }
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "command failed" in full_text

    def test_docx_with_markdown_headers(self, output_dir):
        """Test DOCX renders markdown headers as Word headings."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Explain this"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "text",
                            "text": "# Overview\nSome text\n## Details\nMore text\n### Subsection\nFinal text",
                        }
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Overview" in full_text
        assert "Details" in full_text
        assert "Subsection" in full_text

    def test_docx_with_list_tool_result(self, output_dir):
        """Test DOCX handles tool results that are lists of content blocks."""
        from docx import Document

        loglines = [
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:00.000Z",
                "message": {"role": "user", "content": "Do something"},
            },
            {
                "type": "assistant",
                "timestamp": "2025-01-01T10:00:05.000Z",
                "message": {
                    "role": "assistant",
                    "content": [
                        {
                            "type": "tool_use",
                            "name": "Read",
                            "id": "tool_1",
                            "input": {"file_path": "/tmp/test.py"},
                        }
                    ],
                },
            },
            {
                "type": "user",
                "timestamp": "2025-01-01T10:00:10.000Z",
                "message": {
                    "role": "user",
                    "content": [
                        {
                            "type": "tool_result",
                            "tool_use_id": "tool_1",
                            "content": [
                                {"type": "text", "text": "Line 1 of file"},
                                {"type": "text", "text": "Line 2 of file"},
                            ],
                        }
                    ],
                },
            },
        ]

        docx_path = output_dir / "output.docx"
        generate_docx({"loglines": loglines}, docx_path)

        doc = Document(str(docx_path))
        full_text = "\n".join(p.text for p in doc.paragraphs)
        assert "Line 1 of file" in full_text
        assert "Line 2 of file" in full_text

    def test_docx_cli_format_option(self, output_dir):
        """Test that --format docx option works via CLI."""
        from click.testing import CliRunner
        from claude_code_transcripts import cli

        fixture_path = Path(__file__).parent / "sample_session.json"
        docx_path = output_dir / "output.docx"

        runner = CliRunner()
        result = runner.invoke(
            cli,
            ["json", str(fixture_path), "-o", str(docx_path), "--format", "docx"],
        )

        assert result.exit_code == 0
        assert docx_path.exists()
        assert "Output:" in result.output
